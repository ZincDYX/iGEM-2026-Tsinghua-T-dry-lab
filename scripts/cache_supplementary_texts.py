from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import tempfile
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Callable

import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
SUP_DIR = ROOT / "essay" / "supplementary_current"
DEFAULT_MANIFEST = SUP_DIR / "supplementary_current_download_manifest.csv"
DEFAULT_CACHE_DIR = SUP_DIR / "text_cache"
DEFAULT_OUT = SUP_DIR / "supplementary_current_text_manifest.csv"
DEFAULT_SUMMARY = SUP_DIR / "supplementary_current_text_summary.json"

SUPPORTED_EXT = {".xlsx", ".xls", ".csv", ".tsv", ".txt", ".docx", ".pdf", ".zip", ".bin"}


def norm(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def safe_abs(path_text: str) -> Path:
    path = Path(path_text)
    if not path.is_absolute():
        path = ROOT / path
    return path


def safe_id(row: dict[str, str]) -> str:
    basis = row.get("file_path") or row.get("url") or row.get("doi") or row.get("title") or "supplementary"
    digest = hashlib.sha1(basis.encode("utf-8", errors="ignore")).hexdigest()[:12]
    rank = norm(row.get("rank")) or "0"
    rank_digits = re.sub(r"\D+", "", rank) or "0"
    return f"rank_{int(rank_digits):05d}_{digest}"


def df_to_text(df: pd.DataFrame, label: str, max_rows: int) -> str:
    if df.empty:
        return f"[{label}] empty table"
    df = df.dropna(axis=0, how="all").dropna(axis=1, how="all")
    columns = [norm(c) or f"column_{i+1}" for i, c in enumerate(df.columns)]
    out = [f"[{label}] rows={len(df)} columns={len(columns)}", "columns: " + " | ".join(columns)]
    for idx, (_, row) in enumerate(df.head(max_rows).iterrows(), start=1):
        cells: list[str] = []
        for col, value in zip(columns, row.tolist()):
            text = norm(value)
            if text and text.lower() != "nan":
                cells.append(f"{col}={text}")
        if cells:
            out.append(f"row {idx}: " + "; ".join(cells))
    return "\n".join(out)


def read_excel(path: Path, max_rows: int) -> str:
    parts: list[str] = []
    with pd.ExcelFile(path) as xls:
        for sheet in xls.sheet_names[:30]:
            try:
                df = pd.read_excel(xls, sheet_name=sheet, nrows=max_rows)
                parts.append(df_to_text(df, f"sheet: {sheet}", max_rows=max_rows))
            except Exception as exc:
                parts.append(f"[sheet: {sheet}] error: {type(exc).__name__}: {exc}")
    return "\n\n".join(parts)


def read_delimited(path: Path, max_rows: int) -> str:
    sep = "\t" if path.suffix.lower() == ".tsv" else None
    errors: list[str] = []
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin1"):
        try:
            df = pd.read_csv(path, sep=sep, engine="python", encoding=encoding, nrows=max_rows)
            return df_to_text(df, f"table: {path.name}", max_rows=max_rows)
        except Exception as exc:
            errors.append(f"{encoding}:{type(exc).__name__}")
    raise RuntimeError("csv_parse_failed " + ", ".join(errors[:4]))


def read_plain_text(path: Path, max_chars: int) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin1"):
        try:
            return path.read_text(encoding=encoding, errors="ignore")[:max_chars]
        except Exception:
            continue
    return path.read_bytes()[:max_chars].decode("utf-8", errors="ignore")


def read_docx(path: Path) -> str:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError(f"python-docx unavailable: {exc}") from exc
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table_i, table in enumerate(doc.tables, start=1):
        parts.append(f"[docx table {table_i}]")
        for row_i, row in enumerate(table.rows, start=1):
            cells = [norm(cell.text) for cell in row.cells if norm(cell.text)]
            if cells:
                parts.append(f"row {row_i}: " + " | ".join(cells))
    return "\n".join(parts)


def read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        raise RuntimeError(f"pypdf unavailable: {exc}") from exc
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            parts.append(text.strip())
    return "\n\n".join(parts)


def read_supported_file(path: Path, *, max_rows: int, max_chars: int) -> str:
    ext = path.suffix.lower()
    if ext in {".xlsx", ".xls"}:
        return read_excel(path, max_rows=max_rows)
    if ext in {".csv", ".tsv"}:
        return read_delimited(path, max_rows=max_rows)
    if ext == ".txt":
        return read_plain_text(path, max_chars=max_chars)
    if ext == ".docx":
        return read_docx(path)
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".bin":
        data = path.read_bytes()[:8]
        if data.startswith(b"PK"):
            try:
                return read_excel(path, max_rows=max_rows)
            except Exception:
                return read_zip(path, max_rows=max_rows, max_chars=max_chars)
        return read_plain_text(path, max_chars=max_chars)
    if ext == ".zip":
        return read_zip(path, max_rows=max_rows, max_chars=max_chars)
    raise RuntimeError(f"unsupported extension: {ext}")


def read_zip(path: Path, *, max_rows: int, max_chars: int) -> str:
    parts: list[str] = []
    with zipfile.ZipFile(path) as zf, tempfile.TemporaryDirectory() as td:
        temp_root = Path(td)
        names = [n for n in zf.namelist() if not n.endswith("/")]
        supported = [n for n in names if Path(n).suffix.lower() in SUPPORTED_EXT - {".zip"}]
        for name in supported[:80]:
            target = temp_root / hashlib.sha1(name.encode("utf-8", errors="ignore")).hexdigest()
            target = target.with_suffix(Path(name).suffix)
            target.write_bytes(zf.read(name))
            try:
                text = read_supported_file(target, max_rows=max_rows, max_chars=max_chars)
                if text.strip():
                    parts.append(f"[zip member: {name}]\n{text}")
            except Exception as exc:
                parts.append(f"[zip member: {name}] error: {type(exc).__name__}: {exc}")
    return "\n\n".join(parts)


def existing_files(path: Path, *, retry_errors: bool) -> set[str]:
    if not path.exists():
        return set()
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return {
            r.get("file_path", "")
            for r in csv.DictReader(f)
            if r.get("file_path") and (not retry_errors or r.get("status") == "cached")
        }


def load_rows(path: Path, done: set[str], limit: int) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            file_path = row.get("file_path", "")
            if row.get("status") != "downloaded" or not file_path:
                continue
            if file_path in done or file_path in seen:
                continue
            ext = safe_abs(file_path).suffix.lower()
            if ext not in SUPPORTED_EXT:
                continue
            seen.add(file_path)
            rows.append(row)
            if limit > 0 and len(rows) >= limit:
                break
    return rows


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--cache-dir", type=Path, default=DEFAULT_CACHE_DIR)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--summary", type=Path, default=DEFAULT_SUMMARY)
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--max-rows-per-table", type=int, default=300)
    parser.add_argument("--max-text-chars", type=int, default=220000)
    parser.add_argument("--retry-errors", action="store_true", help="Retry files that already have error/empty rows in the text manifest.")
    args = parser.parse_args()

    args.cache_dir.mkdir(parents=True, exist_ok=True)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    done = existing_files(args.out, retry_errors=args.retry_errors)
    rows = load_rows(args.manifest, done, args.limit)

    fields = [
        "rank", "candidate_id", "doi", "title", "url", "source", "file_path",
        "text_file", "status", "text_chars", "error",
    ]
    write_header = not args.out.exists() or args.out.stat().st_size == 0

    def run_row(row: dict[str, str]) -> dict[str, str]:
        out_row = {field: row.get(field, "") for field in fields}
        out_row["text_file"] = ""
        out_row["status"] = "error"
        out_row["text_chars"] = "0"
        out_row["error"] = ""
        path = safe_abs(row.get("file_path", ""))
        try:
            text = read_supported_file(path, max_rows=max(10, args.max_rows_per_table), max_chars=max(1000, args.max_text_chars))
            header = "\n".join([
                f"doi: {row.get('doi', '')}",
                f"title: {row.get('title', '')}",
                f"supplementary_source: {row.get('url', '')}",
                f"source_file: {path}",
                "",
            ])
            text = (header + text)[: args.max_text_chars]
            text_path = args.cache_dir / f"{safe_id(row)}.txt"
            text_path.write_text(text, encoding="utf-8", errors="ignore")
            out_row["text_file"] = str(text_path)
            out_row["text_chars"] = str(len(text))
            out_row["status"] = "cached" if text.strip() else "empty_text"
        except Exception as exc:
            out_row["error"] = f"{type(exc).__name__}: {exc}"
        return out_row

    status_counts: dict[str, int] = {}
    with args.out.open("a", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        if write_header:
            writer.writeheader()
        with ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
            futures = [pool.submit(run_row, row) for row in rows]
            for i, future in enumerate(as_completed(futures), start=1):
                result = future.result()
                writer.writerow(result)
                f.flush()
                status_counts[result.get("status", "")] = status_counts.get(result.get("status", ""), 0) + 1
                if i % 100 == 0 or i == len(rows):
                    print(f"cached_or_attempted={i}/{len(rows)} status={dict(sorted(status_counts.items()))}", flush=True)

    final_status: dict[str, int] = {}
    final_rows = 0
    cached_dois: set[str] = set()
    if args.out.exists():
        with args.out.open("r", encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                final_rows += 1
                final_status[row.get("status", "")] = final_status.get(row.get("status", ""), 0) + 1
                if row.get("status") == "cached" and row.get("doi"):
                    cached_dois.add(row["doi"].lower())
    summary = {
        "attempted_this_run": len(rows),
        "manifest_rows_total": final_rows,
        "cached_dois_total": len(cached_dois),
        "status_counts": dict(sorted(final_status.items())),
        "text_manifest": str(args.out),
        "cache_dir": str(args.cache_dir),
    }
    args.summary.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

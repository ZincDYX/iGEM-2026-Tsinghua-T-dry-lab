"""Compare the unchanged author RHS with independent published source data.

This reproduces numerical figures and reanalyses supplied measurements, not
genome assembly, MD, wet-lab experiments, or validation of the production model.
All computations are intended to run on air-lab. No fitted parameter is fed
back into the author RHS or the production scientific parameter registry.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.metadata
import importlib.util
import itertools
import json
from pathlib import Path
import platform
import socket
import sys
from types import SimpleNamespace

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from scipy.integrate import solve_ivp
from scipy.optimize import curve_fit
from scipy.stats import kruskal, norm, rankdata

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from igem_drylab.published_reference import load_author_native_parameters

AUTHOR_MD5 = {
    "simulation-network-glabridin-final.py": "43195d5a0ffef38aef774be6e174aa22",
    "simulation-data.txt": "66efd25bea2fe3c4e6d095c44fcb944e",
    "Source Data_V2.xlsx": "ca6192d3a4c11bc1af975ad24d4a0d37",
}


def save_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def save_csv(path, headers, rows):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def figure(path, fig):
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)


def source_table(rows, start, stop):
    """Read a numeric source table, retaining numeric strings and true zeros."""
    selected = []
    for row in rows:
        if row[start] is None or row[start] == "":
            if any(value is not None and value != "" for value in row[start+1:stop]):
                raise ValueError("Source concentrations exist without a time coordinate")
            continue
        # float(None) must fail: a missing concentration is not a measured zero.
        selected.append([float(value) for value in row[start:stop]])
    return np.asarray(selected, dtype=float)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sources", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    raw, out = args.sources.resolve(), args.output.resolve()
    out.mkdir(parents=True, exist_ok=True)
    checks = []
    for name, expected in AUTHOR_MD5.items():
        data = (raw / name).read_bytes()
        actual = hashlib.md5(data).hexdigest()
        if actual != expected:
            raise ValueError(f"Official author MD5 mismatch: {name}")
        checks.append(dict(name=name, md5=actual, sha256=hashlib.sha256(data).hexdigest(), bytes=len(data)))
    save_json(out / "source_checks.json", checks)
    save_json(out / "execution_environment.json", {
        "hostname": socket.gethostname(), "python": sys.version, "platform": platform.platform(),
        "executable": sys.executable, "cwd": str(Path.cwd()),
        "packages": {name: importlib.metadata.version(name) for name in
                     ("numpy", "scipy", "matplotlib", "openpyxl", "requests")},
        "formulas": "unmodified author reaction_kinetics imported directly",
        "solver_default": {"method": "LSODA", "rtol": 1e-3, "atol": 1e-6, "points": 1000},
        "diagnostic_solver": {"method": "LSODA", "rtol": 1e-8, "atol": 1e-12},
    })
    code = raw / "simulation-network-glabridin-final.py"
    spec = importlib.util.spec_from_file_location("immutable_glabridin_author", code)
    author = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(author)
    values = load_author_native_parameters(raw / "simulation-data.txt")
    workbook = load_workbook(raw / "Source Data_V2.xlsx", read_only=True, data_only=True)
    sheets = {sheet.title: list(sheet.values) for sheet in workbook}
    save_json(out / "sheet_inventory.json", [
        {"sheet": name, "rows": len(rows), "columns": len(rows[0]),
         "populated_cells": sum(v is not None for row in rows for v in row), "headers": rows[:3]}
        for name, rows in sheets.items()
    ])
    def find_sheet(label):
        # Exact normalized label: only publisher punctuation/spaces differ.
        key = label.replace(" ", "").replace(".", "")
        matches = [rows for name, rows in sheets.items()
                   if name.replace(" ", "").replace(".", "") == key]
        if len(matches) != 1:
            raise ValueError(f"Ambiguous/missing source sheet: {label}")
        return matches[0]

    def run(ratio, dmt_scale=1.0, tight=False):
        km = [values[f"km{i}"] for i in range(1, 11)]
        kcat = [values[f"kcat{i}dkm{i}"] * km[i-1] for i in range(1, 11)]
        for i in range(4, 7):  # Published note: reactions 5--7, not PTR/PT/OC.
            kcat[i] *= dmt_scale
        rate = values["rate"]
        ratios = [values[f"K{i}"] for i in range(8, 16)]
        params = dict(kcat=kcat, Km=km, E=values["E"],
                      K_out=[(k*rate/(k+1), rate/(k+1)) for k in ratios],
                      disable_times={i: ratio*values["total_time"] for i in range(4, 8)})
        options = dict(rtol=1e-8, atol=1e-12) if tight else {}
        solution = solve_ivp(author.reaction_kinetics, (0, values["total_time"]),
                             [7.4e-5] + [0]*15, args=(params,),
                             t_eval=np.linspace(0, values["total_time"], 1000), method="LSODA", **options)
        if not solution.success:
            raise RuntimeError(solution.message)
        return solution

    def table(solution):
        return np.column_stack([solution.t, (solution.y[:8]+solution.y[8:]).T, solution.y[13]])

    metrics = []
    def compare(name, calculated, source, units):
        if calculated.shape != source.shape:
            raise ValueError(f"{name}: shape mismatch {calculated.shape} != {source.shape}")
        difference = calculated-source
        metrics.append(dict(target=name, shape=list(source.shape), units=units,
                            max_abs_difference=float(np.max(np.abs(difference))),
                            rmse=float(np.sqrt(np.mean(difference**2))),
                            max_abs_time_difference_days=float(np.abs(difference[:,0]).max()),
                            max_abs_concentration_difference=float(np.abs(difference[:,1:]).max()),
                            concentration_rmse=float(np.sqrt(np.mean(difference[:,1:]**2)))))
        save_csv(out / f"{name}_comparison.csv", ["row", "column", "reference", "reproduced", "difference"],
                 ([i, j, source[i,j], calculated[i,j], difference[i,j]]
                  for i,j in itertools.product(range(source.shape[0]), range(source.shape[1]))))

    print("Running unchanged author RHS against published numerical tables", flush=True)
    base = run(values["time_ratio"])
    # Verify the adapter against the actual author's one_step, not our own fixture.
    km = [values[f"km{i}"] for i in range(1, 11)]
    data = SimpleNamespace(kcat_list=[values[f"kcat{i}dkm{i}"]*km[i-1] for i in range(1,11)],
                           km_list=km, K_list=[values[f"K{i}"] for i in range(8,16)],
                           time_ratio=values["time_ratio"], total_time=values["total_time"],
                           E=values["E"], rate=values["rate"])
    author_output = np.asarray(author.one_step(data)).T  # Agg backend: no interactive window.
    plt.close("all")
    compare("author_one_step_adapter", table(base), author_output, "author-native, time days")
    h = source_table(find_sheet("Fig.4h")[3:], 1, 11)
    compare("Fig4h", table(base), h, "source concentration M; time column days")
    g = find_sheet("Fig.4g")
    g_left = source_table(g[3:], 1, 11)
    g_right = source_table(g[3:], 12, 22)
    g_solutions = {}
    # Keep both documented time_ratio=0 and author's existing ungated branch.
    # -1 is not silently substituted: it lies outside the documented range.
    for label, scale, source in [("GgDMT1",1.0,g_left),("NhPDA1",0.1,g_right)]:
        for ratio in [0.0, -1.0]:
            sol = run(ratio,scale)
            name = f"Fig4g_{label}_{'documented_zero' if ratio == 0 else 'undocumented_ungated'}"
            compare(name, table(sol), source, "source concentration M; time days")
            g_solutions[name] = sol
    b = find_sheet("SupplementaryFig.18b")
    b = source_table(b[3:], 1, 8)
    curves = [run(ratio) for ratio in [0,.2,.4,.6,.8,.9]]
    reproduced_b = np.column_stack([base.t, *[(s.y[7]+s.y[15])*1000 for s in curves]])
    compare("SupplementaryFig18b", reproduced_b, b, "concentration mM; time days")
    c = find_sheet("SupplementaryFig.18c")
    c = source_table(c[2:], 1, 3)
    scan, tight_scan = [], []
    for i, (day, reference) in enumerate(c):
        sol = run(day/values["total_time"])
        scan.append([day, (sol.y[7,-1]+sol.y[15,-1])*1000])
        diagnostic_sol = run(day/values["total_time"], tight=True)
        tight_scan.append([day, (diagnostic_sol.y[7,-1]+diagnostic_sol.y[15,-1])*1000])
        if i % 100 == 0:
            print(f"Induction sweep {i}/{len(c)}", flush=True)
    scan = np.asarray(scan)
    tight_scan = np.asarray(tight_scan)
    save_csv(out / "induction_tolerance_sensitivity.csv",["induction_days","author_default_final_mM","tighter_final_mM"],
             zip(scan[:,0],scan[:,1],tight_scan[:,1]))
    compare("SupplementaryFig18c", scan, c, "concentration mM; induction days")
    optimum = int(np.argmax(scan[:,1]))
    tight = run(values["time_ratio"], tight=True)
    diagnostic = dict(min_state_default=float(base.y.min()), min_state_tight=float(tight.y.min()),
                      max_unweighted_sum_drift=float(np.abs(base.y.sum(axis=0)-7.4e-5).max()),
                      max_trajectory_tolerance_change=float(np.abs(base.y-tight.y).max()),
                      final_product_default_M=float(base.y[7,-1]+base.y[15,-1]),
                      final_product_tight_M=float(tight.y[7,-1]+tight.y[15,-1]),
                      sweep_points=len(scan), optimum_grid_day=float(scan[optimum,0]),
                      optimum_grid_final_mM=float(scan[optimum,1]),
                      tight_optimum_grid_day=float(tight_scan[np.argmax(tight_scan[:,1]),0]),
                      tight_optimum_grid_final_mM=float(tight_scan[:,1].max()),
                      note="Sum is the author's equal-weight state sum, NOT volume-weighted amount conservation. No negative values were clipped.")
    save_json(out / "numerical_comparisons.json", metrics)
    save_json(out / "solver_diagnostics.json", diagnostic)
    fig, axs = plt.subplots(2,2,figsize=(12,8))
    for i in range(8):
        axs[0,0].plot(base.t,(base.y[i]+base.y[i+8])*1000,label=f"({i+8})")
    axs[0,0].set(title="Fig.4h unchanged author model",xlabel="Days",ylabel="Concentration (mM)")
    axs[0,0].legend(ncol=4,fontsize=8)
    for day, sol in zip([0,1,2,3,4,4.5],curves):
        axs[0,1].plot(sol.t,(sol.y[7]+sol.y[15])*1000,label=f"Day {day}")
    axs[0,1].set(title="Supplementary Fig.18b",xlabel="Days",ylabel="Glabridin (mM)")
    axs[0,1].legend(fontsize=8)
    axs[1,0].plot(c[:,0],c[:,1],label="Published")
    axs[1,0].plot(scan[:,0],scan[:,1],"--",label="Reproduced")
    axs[1,0].plot(tight_scan[:,0],tight_scan[:,1],":",label="Tighter tolerances (diagnostic)")
    axs[1,0].set(title="Supplementary Fig.18c: full published grid",xlabel="Induction (days)",ylabel="Final glabridin (mM)")
    axs[1,0].legend()
    axs[1,1].plot(base.t,(base.y[7]+base.y[15])*1000,label="Author default tolerances")
    axs[1,1].plot(tight.t,(tight.y[7]+tight.y[15])*1000,label="Tighter diagnostic tolerances")
    axs[1,1].set(title="Numerical sensitivity; same RHS",xlabel="Days",ylabel="Glabridin (mM)")
    axs[1,1].legend(fontsize=8)
    figure(out / "numerical_reproduction.png",fig)
    fig, axes = plt.subplots(1,2,figsize=(12,4))
    for ax, label, source in zip(axes,["GgDMT1","NhPDA1"],[g_left,g_right]):
        for i in [5,7]:
            ax.plot(source[:,0],source[:,i+1]*1000,label=f"Published ({i+8})")
            for suffix, style in [("documented_zero",":"),("undocumented_ungated","--")]:
                sol=g_solutions[f"Fig4g_{label}_{suffix}"]
                ax.plot(sol.t,(sol.y[i]+sol.y[i+8])*1000,style,label=f"{suffix} ({i+8})")
        ax.set(title=f"Fig.4g source label: {label}",xlabel="Days",ylabel="mM")
        ax.legend(fontsize=7)
    figure(out / "fig4g_scenario_audit.png",fig)

    print("Refitting all 12 published kinetic datasets",flush=True)
    kinetics = find_sheet("SupplementaryFig.5")
    fit_rows, observations = [], []
    # Figure annotations read from supplementary PDF page 6, panels a--l.
    paper_km = [32.78,46.54,17.38,186.53,41.68,42.37,30.23,14.27,33.82,21.17,13.814,23.33]
    fig, axes = plt.subplots(4,3,figsize=(13,13))
    for index,(header,end,col) in enumerate((h,e,col) for h,e in [(2,13),(16,27),(30,40),(43,52)] for col in [2,7,12]):
        enzyme, unit = kinetics[header-1][col-1], kinetics[header-1][col]
        substrate, mean, sd, all_x, all_y = [], [], [], [], []
        for rownum in range(header+2,end+1):
            row=kinetics[rownum-1]
            concentration = row[col-1]
            if not isinstance(concentration,(int,float)):
                continue
            reps=[v for v in row[col:col+3] if isinstance(v,(int,float))]
            substrate.append(concentration); mean.append(np.mean(reps)); sd.append(np.std(reps,ddof=1))
            for rep,value in enumerate(row[col:col+3],1):
                if isinstance(value,(int,float)):
                    all_x.append(concentration); all_y.append(value)
                    observations.append([index+1,enzyme,unit,f"{get_column_letter(col+rep)}{rownum}",concentration,value])
        x,y=np.asarray(substrate),np.asarray(mean)
        def mm(s,vmax,km):
            return vmax*s/(km+s)
        fitted,cov=curve_fit(mm,x,y,p0=[max(y),max(x)/10],bounds=(0,np.inf),maxfev=20000,
                            ftol=1e-12,xtol=1e-12,gtol=1e-12)
        all_fit,_=curve_fit(mm,np.asarray(all_x),np.asarray(all_y),p0=fitted,bounds=(0,np.inf),maxfev=20000,
                            ftol=1e-12,xtol=1e-12,gtol=1e-12)
        se=np.sqrt(np.diag(cov))
        fit_rows.append(dict(panel=chr(97+index),enzyme=enzyme,source_velocity_header=unit,
                             concentrations=len(x),measurements=len(all_y),Km_uM=float(fitted[1]),
                             Km_asymptotic_fit_SE_uM=float(se[1]),Vmax_native=float(fitted[0]),
                             Vmax_asymptotic_fit_SE_native=float(se[0]),paper_annotated_Km_uM=paper_km[index],
                             Km_difference_from_figure_uM=float(fitted[1]-paper_km[index]),
                             Km_fit_all_replicates_uM=float(all_fit[1]),
                             R2_point_means=float(1-np.sum((y-mm(x,*fitted))**2)/np.sum((y-y.mean())**2))))
        ax=axes.flat[index]
        ax.errorbar(x,y,yerr=sd,fmt="o",markersize=3,label="Published mean +/- SD")
        grid=np.linspace(0,x.max(),300);ax.plot(grid,mm(grid,*fitted),label="Unweighted MM fit")
        ax.set(title=f"{chr(97+index)} {enzyme}; Km={fitted[1]:.3f} uM",xlabel="Substrate (uM)",ylabel=unit)
    figure(out / "kinetic_fits_12.png",fig)
    save_json(out / "kinetic_fits.json",fit_rows)
    save_csv(out / "kinetic_observations.csv",["dataset","enzyme","velocity_header","source_cell","substrate_uM","velocity_native"],observations)

    print("Reanalysing biological measurements without imputing missing cells",flush=True)
    missing, summaries, wet_points = [],[],[]
    fig,axes=plt.subplots(1,2,figsize=(12,4))
    for ax,label,compounds,columns in [(axes[0],"Fig.4d",[13,9,14,10,15,11],[3,6,9,12,15,18]),
                                      (axes[1],"Fig.4e",[9,13,14,15],[3,6,9,12])]:
        rows=find_sheet(label)
        for compound,col in zip(compounds,columns):
            times,means,sds=[],[],[]
            for rownum,row in enumerate(rows[3:],4):
                if not isinstance(row[1],(int,float)):continue
                reps=[]
                for j in range(col-1,col+2):
                    value=row[j]
                    cell=f"{get_column_letter(j+1)}{rownum}"
                    if isinstance(value,(int,float)):
                        reps.append(value);wet_points.append([label,compound,cell,row[1],value,"mg/L"])
                    else:missing.append(dict(sheet=label,cell=cell,time_h=row[1],compound=compound))
                times.append(row[1]);means.append(np.mean(reps));sds.append(np.std(reps,ddof=1))
                summaries.append(dict(sheet=label,compound=compound,time_h=row[1],n=len(reps),mean_mg_L=float(means[-1]),sd_mg_L=float(sds[-1])))
            ax.errorbar(times,means,yerr=sds,label=f"({compound})",markersize=3,fmt="o-")
        ax.set(title=label+": source data only",xlabel="Time (h)",ylabel="mg/L")
        ax.legend(fontsize=8,ncol=3)
    figure(out / "fig4de_measurements.png",fig)
    f=find_sheet("Fig.4f")
    compounds=[9,13,10,14,11,15]
    groups=[np.asarray([row[col] for row in f[3:] if isinstance(row[col],(int,float))]) for col in range(1,7)]
    transport=[]
    for compound,group in zip(compounds,groups):
        k=values[f"K{compound}"]
        transport.append(dict(compound=compound,n=len(group),mean_percent=float(group.mean()),
                              sd_percent=float(group.std(ddof=1)),median_percent=float(np.median(group)),
                              author_model_equilibrium_percent=k/(1+k)*100))
    stat,p=kruskal(*groups)
    all_values=np.concatenate(groups);ranks=rankdata(all_values)
    means_rank=[];start=0
    for group in groups:
        means_rank.append(np.mean(ranks[start:start+len(group)]));start+=len(group)
    _,ties=np.unique(all_values,return_counts=True);n=len(all_values)
    rank_variance=n*(n+1)/12-np.sum(ties**3-ties)/(12*(n-1))
    dunn=[]
    for i,j in itertools.combinations(range(len(groups)),2):
        z=(means_rank[i]-means_rank[j])/np.sqrt(rank_variance*(1/len(groups[i])+1/len(groups[j])))
        raw_p=float(2*norm.sf(abs(z)))
        dunn.append(dict(compounds=[compounds[i],compounds[j]],z=float(z),p_unadjusted=raw_p,
                         p_bonferroni_sensitivity=min(1.0,raw_p*15)))
    save_json(out / "transport_measurements.json",dict(groups=transport,kruskal_H=float(stat),kruskal_p=float(p),
              dunn=dunn,adjustment_note="Paper names Dunn but does not specify correction; unadjusted and Bonferroni sensitivity are explicitly separate, not an assertion of identical figure p-values."))
    fig,ax=plt.subplots(figsize=(7,4));ax.boxplot(groups,labels=[str(i) for i in compounds],whis=1.5)
    ax.set(title="Fig.4f source extracellular ratios",xlabel="Compound",ylabel="Extracellular (%)")
    figure(out / "fig4f_transport.png",fig)
    strain_rows=[]
    rows=find_sheet("Fig.4i")
    fig,axes=plt.subplots(1,4,figsize=(14,4))
    for ax,compound,col in zip(axes,[2,5,8,15],[3,9,15,21]):
        means,sds=[],[]
        for strain in range(4):
            obs=[row[col-1+strain] for row in rows[3:] if isinstance(row[col-1+strain],(int,float))]
            means.append(np.mean(obs));sds.append(np.std(obs,ddof=1))
            strain_rows.append(dict(compound=compound,strain=strain+1,n=len(obs),mean_mg_L=float(means[-1]),sd_mg_L=float(sds[-1])))
        ax.bar(range(1,5),means,yerr=sds);ax.set(title=f"Fig.4i ({compound})",xlabel="Strain",ylabel="mg/L",xticks=range(1,5))
    figure(out / "fig4i_strains.png",fig)
    save_json(out / "wet_measurements.json",dict(timecourses=summaries,missing_cells=missing,strains=strain_rows))
    save_csv(out / "wet_timecourse_observations.csv",["figure","compound","source_cell","time_h","concentration","unit"],wet_points)
    save_json(out / "reproduction_summary.json",dict(
        numerical=metrics,diagnostics=diagnostic,kinetic_fits=fit_rows,missing_cells=missing,
        completed="Immutable-source verification; original RHS vs independent published tables; complete 500-point induction scan; 12 kinetic fits; Fig4d/e/f/i measurement reanalysis.",
        NOT_completed=["Genome/transcriptome raw-data reprocessing", "Raw metabolomics reconstruction", "Docking and long MD reruns",
                       "Wet-lab experiments", "Validation of condensate transport or production M4"]
    ))
    print(json.dumps(dict(diagnostics=diagnostic,numerical=metrics,missing=missing),indent=2),flush=True)


if __name__ == "__main__":
    main()

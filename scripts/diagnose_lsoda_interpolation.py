"""air-lab-only controlled diagnostic; never patches the production solver."""
import argparse
import ast
from contextlib import contextmanager
import hashlib
import importlib.util
import inspect
import json
from pathlib import Path
import platform
import sys
import zipfile

import matplotlib
matplotlib.use('Agg')
import numpy as np
from openpyxl import load_workbook
import scipy
from scipy.integrate import solve_ivp
import scipy.integrate._ivp.lsoda as lsoda

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from igem_drylab.published_reference import load_author_native_parameters


@contextmanager
def diagnostic_interpolation(method):
    current=lsoda.LSODA._dense_output_impl
    lsoda.LSODA._dense_output_impl=method
    try:
        yield
    finally:
        lsoda.LSODA._dense_output_impl=current


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--sources',type=Path,required=True)
    parser.add_argument('--review-evidence',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    out=args.output.resolve();out.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(args.review_evidence) as archive:
        manifest=json.loads(archive.read('MANIFEST.json'))
        records={item['path']:item for item in manifest['files']}
        source=archive.read('solver_sources/lsoda_v1.10.1.py')
        if hashlib.sha256(source).hexdigest()!=records['solver_sources/lsoda_v1.10.1.py']['sha256']:
            raise ValueError('review_solver_source_hash_mismatch')
        (out/'lsoda_v1.10.1.py').write_bytes(source)
        # Retain the full supplied SciPy license with the diagnostic source.
        (out/'SCIPY_LICENSE.txt').write_bytes(archive.read('solver_sources/SCIPY_LICENSE.txt'))
    klass=next(node for node in ast.parse(source).body if isinstance(node,ast.ClassDef) and node.name=='LSODA')
    function=next(node for node in klass.body if isinstance(node,ast.FunctionDef) and node.name=='_dense_output_impl')
    namespace=dict(lsoda.__dict__)
    exec(compile(ast.Module(body=[function],type_ignores=[]),'official-scipy-v1.10.1-lsoda.py','exec'),namespace)
    modern=lsoda.LSODA._dense_output_impl
    raw=args.sources.resolve()
    code=raw/'simulation-network-glabridin-final.py'
    spec=importlib.util.spec_from_file_location('immutable_author_lsoda_diagnostic',code)
    author=importlib.util.module_from_spec(spec);spec.loader.exec_module(author)
    values=load_author_native_parameters(raw/'simulation-data.txt')
    workbook=load_workbook(raw/'Source Data_V2.xlsx',read_only=True,data_only=True)
    def table(label,start,stop):
        sheet=next(workbook[name] for name in workbook.sheetnames if name.replace(' ','').replace('.','')==label)
        return np.asarray([[float(v) for v in row[start:stop]] for row in list(sheet.values)[3:] if row[start] is not None])
    cases=[('Fig4h',values['time_ratio'],1,table('Fig4h',1,11)),
           ('Fig4g_GgDMT1_ungated',-1,1,table('Fig4g',1,11)),
           ('Fig4g_NhPDA1_ungated',-1,.1,table('Fig4g',12,22))]
    workbook.close()
    report={'execution_environment':{'python':sys.version,'scipy':scipy.__version__,'numpy':np.__version__,'platform':platform.platform()},
        'controlled_variable':'Only LSODA dense-output method; same modern integration backend, unchanged author RHS/parameters/tolerances',
        'official_sources':['https://github.com/scipy/scipy/issues/14480',
            'https://docs.scipy.org/doc/scipy-1.12.0/release/1.12.0-notes.html',
            'https://raw.githubusercontent.com/scipy/scipy/v1.10.1/scipy/integrate/_ivp/lsoda.py'],
        'old_solver_source_sha256':hashlib.sha256(source).hexdigest(),
        'current_dense_output_source_sha256':hashlib.sha256(inspect.getsource(modern).encode()).hexdigest(),
        'rhs_sha256':hashlib.sha256(code.read_bytes()).hexdigest(),
        'author_parameters_sha256':hashlib.sha256((raw/'simulation-data.txt').read_bytes()).hexdigest(),
        'source_file_sha256':hashlib.sha256((raw/'Source Data_V2.xlsx').read_bytes()).hexdigest(),
        'solver_settings':{'method':'LSODA','rtol':1e-3,'atol':1e-6,'points':1000},'results':[]}
    solutions={}
    def parameters(ratio,scale):
        km=[values[f'km{i}'] for i in range(1,11)]
        kcat=[values[f'kcat{i}dkm{i}']*km[i-1] for i in range(1,11)]
        for index in (4,5,6):kcat[index]*=scale
        return {'kcat':kcat,'Km':km,'E':values['E'],
            'disable_times':{i:ratio*values['total_time'] for i in range(4,8)},
            'K_out':[(values[f'K{i}']*values['rate']/(values[f'K{i}']+1),values['rate']/(values[f'K{i}']+1)) for i in range(8,16)]}
    for mode,method in [('current',modern),('official_v1.10.1_dense_output',namespace['_dense_output_impl'])]:
        with diagnostic_interpolation(method):
            for name,ratio,scale,reference in cases:
                params=parameters(ratio,scale)
                solved=solve_ivp(author.reaction_kinetics,(0,values['total_time']),[7.4e-5]+[0.]*15,
                    args=(params,),t_eval=np.linspace(0,values['total_time'],1000),method='LSODA',rtol=1e-3,atol=1e-6)
                if not solved.success:raise RuntimeError(solved.message)
                calculated=np.column_stack([solved.t,(solved.y[:8]+solved.y[8:]).T,solved.y[13]])
                report['results'].append({'interpolation':mode,'target_id':name,'scenario':{'time_ratio':ratio,'R5_R6_R7_kcat_scale':scale},
                    'max_concentration_difference':float(np.max(np.abs(calculated[:,1:]-reference[:,1:]))),
                    'final15':float(solved.y[7,-1]+solved.y[15,-1]),'min_raw_state':float(solved.y.min())})
                solutions[mode,name]=solved.y
                np.savez_compressed(out/f'{name}_{mode}_raw_states.npz',time=solved.t,state=solved.y)
    report['endpoint_invariance']={name:bool(np.array_equal(solutions['current',name][:,-1],solutions['official_v1.10.1_dense_output',name][:,-1])) for name,*_ in cases}
    report['production_solver_not_patched']=lsoda.LSODA._dense_output_impl is modern
    if not all(report['endpoint_invariance'].values()) or not report['production_solver_not_patched']:
        raise AssertionError('controlled_diagnostic_invariant_failed')
    (out/'lsoda_interpolation_diagnosis.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))
    # Independent numerical convergence on the SAME author RHS. Segment at all
    # induction starts and the author's one-day ramp caps; do not clip states.
    params=parameters(values['time_ratio'],1.)
    end=values['total_time']
    boundaries=sorted({0.,end,*[t for start in params['disable_times'].values() for t in (start,start+1.) if 0<t<end]})
    time=np.unique(np.r_[np.linspace(0,end,1001),boundaries])
    convergence=[];states={}
    for method in ('LSODA','BDF','Radau'):
        for rtol,atol in ((1e-8,1e-12),(1e-10,1e-14),(1e-12,1e-16)):
            initial=np.array([7.4e-5]+[0.]*15)
            state=np.zeros((16,len(time)))
            for left,right in zip(boundaries,boundaries[1:]):
                solved=solve_ivp(author.reaction_kinetics,(left,right),initial,args=(params,),method=method,rtol=rtol,atol=atol,dense_output=True)
                if not solved.success:raise RuntimeError(solved.message)
                selected=(time>=left)&(time<=right)
                state[:,selected]=solved.sol(time[selected])
                initial=solved.y[:,-1]
            key=f'{method}_{rtol}_{atol}'
            states[key]=state
            np.savez_compressed(out/f'convergence_{key}_raw_states.npz',time=time,state=state)
            convergence.append({'method':method,'rtol':rtol,'atol':atol,'boundaries_days':boundaries,
                'minimum_raw_state':float(state.min()),'final15':float(state[7,-1]+state[15,-1]),
                'unweighted_backbone_max_drift':float(np.max(np.abs(state.sum(axis=0)-7.4e-5)))})
    reference=states['Radau_1e-12_1e-16']
    for record in convergence:
        state=states[f"{record['method']}_{record['rtol']}_{record['atol']}"]
        record['max_state_difference_to_tightest_radau']=float(np.max(np.abs(state-reference)))
        record['endpoint15_difference_to_tightest_radau']=float(state[7,-1]+state[15,-1]-reference[7,-1]-reference[15,-1])
    (out/'author_rhs_convergence.json').write_text(json.dumps({'rhs_sha256':report['rhs_sha256'],
        'parameter_file_sha256':report['author_parameters_sha256'],'scenario':{'time_ratio':values['time_ratio'],'R5_R6_R7_kcat_scale':1},
        'raw_states_never_clipped':True,'results':convergence,
        'scope':'numerical convergence diagnostic of author model, not independent biological validation'},indent=2,allow_nan=False),encoding='utf-8')
    print(json.dumps({'convergence':convergence},indent=2))


if __name__=='__main__':main()

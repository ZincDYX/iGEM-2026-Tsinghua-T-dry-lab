from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ESSAY = ROOT / "essay"
SUMMARY = ESSAY / "known_clean_showcase_summary.json"
OUT_DOCX = ESSAY / "全部图表中文说明_Nature风格图集解读.docx"


ACCENT = "7D95A8"
ACCENT2 = "9AAB8F"
LIGHT = "F3F1EE"
HEADER = "E8E3DD"
INK = "222222"
MUTED = "66615E"


MAIN_FIGURES = [
    ("Fig. 1", "Traceable database construction", "展示从全部候选记录到可溯源数值记录、known-clean 记录、core analysis 记录的筛选流程。", "用于说明数据库不是简单堆积 LLM 输出，而是保留 DOI、标题、证据句，并按可追溯性与字段完整度分层。"),
    ("Fig. 2", "Known-field coverage", "展示各关键字段的已知率，包括产物、宿主、反应体系、温度、pH、酶名、酶来源等。", "用于说明哪些字段已经足够可靠，哪些字段因文献缺失而保留 unknown，体现质量控制而非强行补全。"),
    ("Fig. 3a", "Product families", "展示已知产物家族的分布，突出 flavonoid、terpenoid、isoprenoid、LLPS system、isoflavonoid 等类别。", "可用于支持天然产物研究热点集中在类黄酮/萜类，而异黄酮记录相对较少，从而支撑项目创新性。"),
    ("Fig. 3b", "Expression host groups", "展示宿主大类分布，包括 yeast、bacteria、plant、mammalian、cell-free 等。", "用于说明酵母和细菌是主流工程宿主，为后续选择 S. cerevisiae 或 E. coli 对照体系提供文献依据。"),
    ("Fig. 3c", "Reaction system types", "区分 whole-cell、in vivo、fermentation、cell-free、purified enzyme 等实验体系。", "避免将发酵产量记录和体外酶学记录混在一起解释，有利于比较 LLPS 对细胞内与体外体系的不同影响。"),
    ("Fig. 4a", "Quality-control levels", "展示 A_complete、B_product-condition、C_partial-context、D_result-only 等质量层级。", "用于说明数据库既有严格记录层，也保留高召回证据层；展示结论主要基于 known-clean/core 层。"),
    ("Fig. 4b", "Database tiers", "展示 core、LLPS core、silver、review-level 等数据库层级。", "用于说明项目将强相关实验、LLPS 核心记录、综述级背景证据分开管理。"),
    ("Fig. 4c", "Evidence and completion methods", "展示字段来源与补全方法，包括 field_context_search、doi_consensus、doi_consensus_after_pro、rule_lexicon、supplementary_consensus 等。", "这是方法学贡献的核心图：说明数据库经过动态 ICL、全文复核、SI/data 共识和 DOI 共识多轮质控。"),
    ("Fig. 5a", "Titer distribution", "以 log10(mg/L) 展示滴度跨度，并标注项目起点 0.5 mg/L 与目标 5 mg/L。", "用于说明天然产物异源合成滴度跨度很大，项目目标处于文献分布中可讨论的合理区间。"),
    ("Fig. 5b", "Temperature distribution", "展示已知反应或培养温度分布。", "可用于给出 LLPS-酶催化体系的温度参考范围，常见微生物工程条件集中在 30-37 deg C 附近。"),
    ("Fig. 5c", "pH distribution", "展示已知反应 pH 分布。", "用于说明多数记录接近中性条件，为设计酶活与 LLPS 兼容实验提供条件参考。"),
    ("Fig. 5d", "Reaction time distribution", "以小时为单位展示反应/发酵时间跨度。", "用于区分短时间体外酶活实验和长时间细胞工厂生产实验。"),
    ("Fig. 6", "Dynamic ICL and multi-pass QC workflow", "展示文献池、动态 ICL、LLM 抽取、规则约束、全文/SI 质控、可追溯数据库的流程。", "这是项目 dry-lab 方法主图，可直接放入 PPT 说明 AI Agent 如何从文献构建结构化知识库。"),
]

V1_PAGES = [
    ("Appendix Fig. A.1", "Field coverage and product-family overview", "展示全文补全后字段覆盖率、产物家族分布和宿主分布。", "用于从旧版图集中回看字段完整度和天然产物家族格局。"),
    ("Appendix Fig. A.2", "Source mode, database tier and metric distribution", "展示来源模式、数据库 tier 和主要结果指标分布。", "用于说明数据来源结构和不同指标类型的组成。"),
    ("Appendix Fig. A.3", "Product-family by host matrix and field co-occurrence", "展示产物家族与宿主矩阵、已知字段共现关系。", "用于识别哪些产物-宿主组合更常见，以及哪些字段经常同时被报道。"),
    ("Appendix Fig. A.4", "Source-product-host-metric flow and network", "展示数据来源到产物家族、宿主类型、指标类型的流向，并包含关系网络。", "用于说明证据链如何连接来源、体系和结果。"),
    ("Appendix Fig. A.5", "Treemap and word-cloud style summaries", "展示数据库 tier treemap、产物家族 treemap 和高频词。", "用于快速呈现数据结构和文本主题。"),
    ("Appendix Fig. A.6", "Reaction condition distributions", "展示温度、pH、反应时间和体系类型分布。", "用于归纳常见实验条件。"),
    ("Appendix Fig. A.7", "Titer benchmark and product-family titer comparison", "展示滴度 log 分布、产品家族箱线/分布比较及关键 benchmark。", "用于将项目起点和目标与文献滴度范围比较。"),
    ("Appendix Fig. A.8", "OC/DMT-like enzymatic parameter landscape", "展示 Km、specific activity、kcat/Km 等酶学参数分布。", "用于辅助讨论底物识别、酶活性和催化效率。"),
    ("Appendix Fig. A.9", "LLPS catalytic effect and trigger classes", "展示 LLPS 相关催化提升倍数和触发方式分类。", "用于支持 LLPS 调控具有文献报道的催化增益与多样触发机制。"),
]

V2_PAGES = [
    ("Appendix Fig. B.1", "Refined field coverage and category overview", "新版用更紧凑版式重绘字段覆盖率、产物和宿主概览。", "适合展示字段覆盖率和主要类别时使用，版面更接近 Nature panel。"),
    ("Appendix Fig. B.2", "Refined source mode, tier and metric distribution", "新版重绘来源、tier 和结果指标分布。", "更适合在方法学部分说明数据库层级和指标构成。"),
    ("Appendix Fig. B.3", "Matrix, evidence chain and comparison plots", "新版整合产物-宿主矩阵、证据链和研究比较图。", "强调可溯源证据链和研究比较，是展示数据库构建质量的关键页。"),
    ("Appendix Fig. B.4", "Source-product-host-metric flow", "新版保留 source-product-host-metric flow。", "用于说明不同来源和记录类型如何汇入最终分析指标。"),
    ("Appendix Fig. B.5", "Top terms and ranked product/host summary", "新版突出高频词和产品/宿主排名。", "适合快速说明文献语料的主题集中区域。"),
    ("Appendix Fig. B.6", "Refined reaction condition distributions", "新版展示温度、pH、时间和体系类型分布。", "可直接支持实验条件选择和 assay reference range。"),
    ("Appendix Fig. B.7", "Refined titer benchmark", "新版展示总体滴度分布、产品家族滴度比较和关键 benchmark 数值。", "适合项目展示中说明 0.5 mg/L 起点与 5 mg/L 目标的合理性。"),
    ("Appendix Fig. B.8", "Refined enzymatic parameter landscape", "新版展示 Km、specific activity、kcat/Km 和项目逻辑说明。", "用于从酶学角度解释底物触发 SBD 与底物被催化的关系。"),
    ("Appendix Fig. B.9", "Refined LLPS effect and trigger classes", "新版展示 LLPS fold-change、trigger classes 和调控类型。", "用于支撑 LLPS 作为调控手段确有文献基础。"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size: int | float = 10.5, bold: bool = False, color: str = INK) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    set_run_font(run, 10.5)
    p.paragraph_format.space_after = Pt(6)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run, 10.5)
    p.paragraph_format.space_after = Pt(3)


def add_table(doc: Document, rows: list[tuple[str, str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [Cm(3.2), Cm(4.2), Cm(7.2), Cm(7.2)]
    headers = ["图号", "图名", "包含的信息", "图意与展示用途"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, text, bold=True)
        set_cell_shading(cell, HEADER)
        cell.width = widths[i]
    for fig_id, title, info, meaning in rows:
        cells = table.add_row().cells
        for i, text in enumerate([fig_id, title, info, meaning]):
            set_cell_text(cells[i], text)
            cells[i].width = widths[i]
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(level=level)
    r = p.add_run(text)
    set_run_font(r, 16 if level == 1 else 13, bold=True, color=ACCENT if level == 1 else "4F6F7C")
    # Remove the empty run Word may create for heading paragraphs.
    if p.runs and not p.runs[0].text:
        p._p.remove(p.runs[0]._r)


def main() -> int:
    summary = json.loads(SUMMARY.read_text(encoding="utf-8")) if SUMMARY.exists() else {}
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("全部图表中文说明：可溯源 LLPS-酶催化文献数据库图集")
    set_run_font(run, 22, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("用于 iGEM 项目展示、PPT 图注整理与 dry-lab 方法学说明")
    set_run_font(run, 11, color=MUTED)

    add_heading(doc, "1. 图集总览", 1)
    add_paragraph(
        doc,
        f"本图集整合了 known-clean 展示报告中的主图，以及两个此前生成的 Nature-style PDF 图表版本。"
        f"当前数据库包含 {summary.get('all_rows', '20,567')} 条候选结构化记录，其中 known-clean 展示层为 {summary.get('known_clean_rows', '19,032')} 条，"
        f"core analysis 严格分析层为 {summary.get('core_known_rows', '8,417')} 条，覆盖 {summary.get('known_clean_unique_dois', '1,339')} 个可追溯 DOI。"
    )
    add_paragraph(
        doc,
        "说明文档的重点不是重复图中每一个数值，而是解释每张图回答的问题、包含的信息以及它在项目展示中的作用。"
    )

    add_heading(doc, "2. 方法学贡献可以如何表述", 1)
    add_bullet(doc, "可溯源数据库：每条结构化记录保留 DOI、标题、数值、单位和证据句，图表结论可以回到原文核查。")
    add_bullet(doc, "动态 ICL：抽取时不是固定五个示例，而是根据文献片段、字段类型和领域语义动态匹配最相关的上下文学习材料。")
    add_bullet(doc, "多轮质控：结合 JSON Schema、单位/范围/非零约束、全文 Pro 复核、SI/data 共识和 DOI 级保守传播，逐步减少错误和 unknown。")
    add_bullet(doc, "分层使用数据：展示层使用 known-clean 记录，建模或严格结论优先使用 core analysis 记录，高召回候选继续作为证据池保留。")

    add_heading(doc, "3. 主报告图逐图说明", 1)
    add_table(doc, MAIN_FIGURES)

    add_heading(doc, "4. 旧版 PDF 图集 version 1 逐页说明", 1)
    add_paragraph(doc, "Version 1 是较早的 Nature-style 完整图集，包含字段覆盖、来源、网络、treemap、条件分布、滴度 benchmark、酶学参数和 LLPS 效应等页面。")
    add_table(doc, V1_PAGES)

    add_heading(doc, "5. 旧版 PDF 图集 version 2 逐页说明", 1)
    add_paragraph(doc, "Version 2 在 version 1 基础上进一步压缩版式、优化 panel 编号和图例空间，更适合作为展示备选图集。")
    add_table(doc, V2_PAGES)

    add_heading(doc, "6. PPT 展示时可以强调的结论", 1)
    add_bullet(doc, "我们不是只做了一个文献列表，而是构建了带 DOI、证据句和质量分层的可溯源结构化数据库。")
    add_bullet(doc, "天然产物数据库显示 flavonoid、terpenoid、isoprenoid 是研究热点，而 isoflavonoid 记录相对稀少，支持本项目切入点的新颖性。")
    add_bullet(doc, "宿主分布显示 yeast 和 bacteria 是主流工程体系，项目选择酵母体系具有充分文献基础。")
    add_bullet(doc, "温度、pH、反应时间和体系类型图可以为后续湿实验设置提供参考范围。")
    add_bullet(doc, "LLPS 图说明已有文献报道相分离对催化或空间组织有提升作用，也整理了常见触发方式，为底物诱导相分离设计提供依据。")
    add_bullet(doc, "动态 ICL + 多轮质控是 dry-lab 的核心技术贡献：它使数据库能够迁移到新体系，而不依赖人工固定 prompt 和少量示例。")

    doc.save(OUT_DOCX)
    print(f"Wrote: {OUT_DOCX}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
